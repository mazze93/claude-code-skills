"""Data tabs: one Excel Table per ledger table, with icons, chips, cross-links.

Every data tab has the same anatomy (so it reads the same everywhere):
  rows 1–2  navy title band, ⌂ Index link
  row  3    one-line "how to read this tab"
  row  4    header (Excel Table header, filter buttons)
  row  5+   data, frozen under the header
Side effects: adds sheets to the workbook passed in.
"""
from openpyxl.formatting.rule import FormulaRule, IconSetRule, DataBarRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from .. import taxonomy as T
from . import styles as S

HEAD, FIRST = 4, 5


def col_letter(spec: list, header: str) -> str:
    return get_column_letter([h for h, *_ in spec].index(header) + 1)


def write_table(wb, name, title, how_to_read, spec, rows, group, table_name):
    """spec: [(header, getter(row)->value, width, styler(cell,row)|None)]. Returns the sheet."""
    ws = wb.create_sheet(name)
    ws.sheet_properties.tabColor = S.TAB_GROUP[group]
    S.band(ws, title, f"{len(rows)} rows · filter with the header arrows · ⌂ returns to the index", len(spec))
    ws.cell(row=3, column=2, value=how_to_read).font = S.font(9, italic=True, color=S.INK_2)
    S.header_row(ws, HEAD, [h for h, *_ in spec])
    for i, r in enumerate(rows):
        for c, (_h, get, _w, styler) in enumerate(spec, start=1):
            cell = ws.cell(row=FIRST + i, column=c, value=get(r))
            cell.font, cell.alignment, cell.border = S.font(10), S.WRAP, S.BOX
            if styler:
                styler(cell, r)
    S.set_widths(ws, [w for _h, _g, w, _s in spec])
    ws.freeze_panes = ws.cell(row=FIRST, column=3)
    if rows:
        ref = f"A{HEAD}:{get_column_letter(len(spec))}{FIRST + len(rows) - 1}"
        tbl = Table(displayName=table_name, ref=ref)
        tbl.tableStyleInfo = TableStyleInfo(name="TableStyleLight1", showRowStripes=False)
        ws.add_table(tbl)
    ws.page_setup.orientation, ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = "landscape", 1, 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True
    return ws


def tint_rows_by(ws, spec, header, colours: dict, n_rows: int) -> None:
    """Tint whole rows by the value in one column (conditional, so edits re-colour)."""
    col = col_letter(spec, header)
    span = f"A{FIRST}:{get_column_letter(len(spec))}{FIRST + max(n_rows, 1) + 200}"
    for value, hex_ in colours.items():
        ws.conditional_formatting.add(span, FormulaRule(formula=[f'${col}{FIRST}="{value}"'],
                                                        fill=S.fill(S.tint(hex_, 0.86))))


# ── Stylers ──────────────────────────────────────────────────────────────
def link_to(sheet: str, rows_by_key: dict, key_field: str):
    def styler(cell, r):
        row = rows_by_key.get(r.get(key_field))
        if row:
            cell.hyperlink, cell.font = f"#'{sheet}'!A{row}", S.LINK
    return styler


def url_link(cell, r):
    if isinstance(cell.value, str) and cell.value.startswith("http"):
        cell.hyperlink, cell.font = cell.value, S.LINK


def chip_by(colour_map: dict, field: str):
    return lambda cell, r: S.chip(cell, colour_map.get(r.get(field), S.MUTED))


def centred(cell, _r):
    cell.alignment = S.CENTER
    cell.font = S.font(12)


# ── Tabs ─────────────────────────────────────────────────────────────────
def calls_tab(wb, calls, pos):
    spec = [
        ("Call", lambda r: r["call_id"], 11, None),
        ("Exchange", lambda r: r["exchange"], 9, link_to("Exchanges", pos["exchange"], "exchange")),
        ("", lambda r: T.OUTCOME_ICON[r["outcome"]], 4, centred),
        ("Outcome", lambda r: r["outcome"], 15, chip_by(T.OUTCOME_COLOUR, "outcome")),
        ("How known", lambda r: r["outcome_source"], 15, None),
        ("Family", lambda r: f"{T.FAMILY_ICON[r['family']]} {r['family']}", 16, chip_by(T.FAMILY_COLOUR, "family")),
        ("Tool", lambda r: r["tool_short"], 20, None),
        ("What it was for", lambda r: r["description"], 40, None),
        ("Input (abridged)", lambda r: r["input"][:240], 50, None),
        ("Seconds", lambda r: r["duration_s"], 9, None),
        ("Error signature", lambda r: r["error_signature"], 16, None),
        ("Error line", lambda r: r["error_excerpt"], 44, None),
        ("Recorded by", lambda r: r["record_source"], 11, None),
        ("Failure", lambda r: r["failure_id"], 9, link_to("Failures", pos["failure"], "failure_id")),
    ]
    ws = write_table(wb, "Tool Calls", "🧾 Tool calls — every call, with what is known about its result",
                     "Outcome says what happened; 'How known' says how we know. ❔ means the source kept the call "
                     "but not its result — that gap is data too.", spec, calls, "ledger", "tblCalls")
    tint_rows_by(ws, spec, "Outcome", {k: v for k, v in T.OUTCOME_COLOUR.items() if k != "ok" and k != "unrecorded"}, len(calls))
    sec = col_letter(spec, "Seconds")
    ws.conditional_formatting.add(f"{sec}{FIRST}:{sec}{FIRST + len(calls)}",
                                  DataBarRule(start_type="num", start_value=0, end_type="max", color="86B6EF"))
    return spec


def failures_tab(wb, failures, pos):
    spec = [
        ("Failure", lambda r: r["failure_id"], 9, None),
        ("Sev", lambda r: r["severity"], 6, centred),
        ("Layer", lambda r: f"{T.LAYER_ICON[r['layer']]} {r['layer']}", 15, chip_by(T.LAYER_COLOUR, "layer")),
        ("Class", lambda r: r["class"], 18, None),
        ("What happened", lambda r: r["summary"], 54, None),
        ("Caught by", lambda r: f"{T.DETECTION_ICON[r['detection']]} {r['detection']}", 16, None),
        ("Who", lambda r: r["detected_by"], 8, None),
        ("When", lambda r: f"{T.CAUGHT_ICON[r['caught']]} {r['caught']}", 17, None),
        ("Happened in", lambda r: r["exchange"], 10, link_to("Exchanges", pos["exchange"], "exchange")),
        ("Caught in", lambda r: r["detected_in"], 9, None),
        ("Exchanges to catch", lambda r: r["exchanges_to_detect"], 10, centred),
        ("Call", lambda r: r["call_id"], 11, link_to("Tool Calls", pos["call"], "call_id")),
        ("Link basis", lambda r: r["link_basis"], 15, None),
        ("Status", lambda r: r["status"], 11, None),
        ("Evidence (verbatim)", lambda r: r.get("evidence", ""), 46, None),
    ]
    ws = write_table(wb, "Failures", "🚨 Failures — what went wrong, what caught it, how late",
                     "Severity: 1 noise · 2 wrong output, caught · 3 would ship (or shipped) a false claim or broken page. "
                     "Link basis says how firmly the failure is tied to a call.", spec, failures, "ledger", "tblFailures")
    sev = col_letter(spec, "Sev")
    ws.conditional_formatting.add(f"{sev}{FIRST}:{sev}{FIRST + len(failures) + 200}",
                                  IconSetRule("3Flags", "num", [1, 2, 3], showValue=True, reverse=True))
    tint_rows_by(ws, spec, "When", {f"{T.CAUGHT_ICON['after-delivery']} after-delivery": "D03B3B"}, len(failures))
    return spec
