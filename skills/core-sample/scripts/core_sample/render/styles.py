"""Visual system for the core-sample workbook.

Chrome: Kintsugi navy + gold (Mazze's site palette). Data colours come from
taxonomy.py (dataviz reference palette for categories, status palette for
outcomes) so charts and cells always agree. Pure definitions + small helpers
that style cells in place.
"""
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter

FONT = "Arial"
NAVY, NAVY_2, GOLD, INK, INK_2, MUTED = "1F2A44", "2C3A5C", "CDA24E", "0B0B0B", "52514E", "898781"
PAPER, HAIRLINE, BASELINE = "FCFCFB", "E1E0D9", "C3C2B7"

# Tab colours by group: the colour of a tab tells you what kind of page it is.
TAB_GROUP = {"overview": NAVY, "ledger": "2A78D6", "work": "1BAF7A", "action": "D03B3B", "reference": MUTED}


def font(size=10, bold=False, color=INK, italic=False) -> Font:
    return Font(name=FONT, size=size, bold=bold, color=color, italic=italic)


def fill(hex_: str) -> PatternFill:
    return PatternFill("solid", fgColor=hex_)


def tint(hex_: str, amount: float = 0.82) -> str:
    """Mix a colour toward white: row tints that keep text at full contrast."""
    r, g, b = (int(hex_[i:i + 2], 16) for i in (0, 2, 4))
    mix = lambda c: round(c + (255 - c) * amount)
    return f"{mix(r):02X}{mix(g):02X}{mix(b):02X}"


HAIR = Side(style="thin", color=HAIRLINE)
BOX = Border(left=HAIR, right=HAIR, top=HAIR, bottom=HAIR)
WRAP = Alignment(wrap_text=True, vertical="top")
CENTER = Alignment(horizontal="center", vertical="center", wrap_text=True)
LEFT_MID = Alignment(horizontal="left", vertical="center", wrap_text=True)

H_TITLE = font(18, True, "FFFFFF")
H_SUB = font(10, False, GOLD, italic=True)
H_SECTION = font(12, True, NAVY)
LINK = Font(name=FONT, size=10, color="1C5CAB", underline="single")
HEAD_FONT = font(10, True, "FFFFFF")
HEAD_FILL = fill(NAVY)


def band(ws, title: str, subtitle: str, width_cols: int, back_link: bool = True) -> None:
    """Navy title band across rows 1–2, with a ⌂ link back to the index in A1."""
    for r in (1, 2):
        for c in range(1, width_cols + 1):
            ws.cell(row=r, column=c).fill = fill(NAVY)
    if back_link:
        home = ws.cell(row=1, column=1, value="⌂ Index")
        home.hyperlink, home.font, home.alignment = "#'Index'!A1", font(10, True, GOLD), LEFT_MID
    ws.cell(row=1, column=2, value=title).font = H_TITLE
    ws.cell(row=2, column=2, value=subtitle).font = H_SUB
    ws.row_dimensions[1].height, ws.row_dimensions[2].height = 30, 18


def header_row(ws, row: int, headers: list[str]) -> None:
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.font, cell.fill, cell.alignment, cell.border = HEAD_FONT, HEAD_FILL, CENTER, BOX
    ws.row_dimensions[row].height = 30


def set_widths(ws, widths: list[float]) -> None:
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w


def chip(cell, hex_: str, text_color: str = INK) -> None:
    """A coloured tag cell: tinted fill, coloured left rule, ink text (never colour-only)."""
    cell.fill = fill(tint(hex_, 0.70))
    cell.font = font(10, True, text_color)
    cell.border = Border(left=Side(style="thick", color=hex_), top=HAIR, bottom=HAIR, right=HAIR)
    cell.alignment = LEFT_MID
