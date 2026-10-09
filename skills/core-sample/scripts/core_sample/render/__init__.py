"""render_workbook(tables, out_path): the core-sample workbook from a ledger.

Order of work: data tabs first (so row positions and column letters exist),
then Chart Data + Charts, then the Index that points at all of it.
Side effect: writes one .xlsx file.
"""
from pathlib import Path

from openpyxl import Workbook

from .. import taxonomy as T
from . import analysis as A
from . import charts as C
from . import styles as S
from .index import build_dictionary, build_index
from .tables import FIRST, calls_tab, col_letter, failures_tab
from .tabs_more import deliverables_tab, events_tab, exchanges_tab, findings_tab, open_items_tab, sources_tab

ORDER = ["Index", "Charts", "Exchanges", "Tool Calls", "Failures", "Findings", "Sources", "Deliverables",
         "Open Items", "Events", "Dictionary", "Chart Data"]


def positions(tables) -> dict:
    """Row numbers for cross-links, computed before anything is written."""
    calls = tables["tool_calls"]
    first_call = {}
    for i, c in enumerate(calls):
        first_call.setdefault(c["exchange"], FIRST + i)
    return {"call": {c["call_id"]: FIRST + i for i, c in enumerate(calls)},
            "failure": {f["failure_id"]: FIRST + i for i, f in enumerate(tables["failures"])},
            "exchange": {e["exchange"]: FIRST + i for i, e in enumerate(tables["exchanges"])},
            "first_call": first_call}


def column_map(calls_spec, failures_spec) -> dict:
    c, f = (lambda h: col_letter(calls_spec, h)), (lambda h: col_letter(failures_spec, h))
    return {"c_id": c("Call"), "c_ex": c("Exchange"), "c_out": c("Outcome"), "c_fam": c("Family"),
            "c_tool": c("Tool"), "c_sec": c("Seconds"), "f_id": f("Failure"), "f_layer": f("Layer"),
            "f_det": f("Caught by"), "f_when": f("When"), "f_ttd": f("Exchanges to catch"), "f_ex": f("Happened in")}


def build_charts(wb, tables, cols):
    """Chart Data blocks + Charts tab. Returns {question key: takeaway} for the index."""
    data_ws, charts_ws = wb.create_sheet("Chart Data"), wb["Charts"]
    S.band(data_ws, "🧮 Chart Data — the numbers behind every chart", "Live formulas over the data tabs; this is "
           "each chart's table view.", 8)
    data_ws.sheet_properties.tabColor = S.TAB_GROUP["reference"]
    S.set_widths(data_ws, [22] + [15] * 8)
    S.band(charts_ws, "📊 Charts — what the ledger says", "Each chart: the question, a one-line answer generated "
           "from the data, then the chart. Numbers are in Chart Data.", 16)
    charts_ws.sheet_properties.tabColor = S.TAB_GROUP["overview"]
    charts_ws.sheet_view.showGridLines = False
    charts_ws.page_setup.orientation, charts_ws.page_setup.fitToWidth, charts_ws.page_setup.fitToHeight = "landscape", 1, 0
    charts_ws.sheet_properties.pageSetUpPr.fitToPage = True
    B = C.Blocks(data_ws)
    fam_c, out_c = [f[1] for f in T.FAMILIES], [o[1] for o in T.OUTCOMES]
    plan = [
        ("work", "Where did the work happen?", A.q_work_by_exchange(B, cols, tables), fam_c, {}),
        ("visibility", "How much of the tool record can anyone check?", A.q_visibility(B, cols, tables), out_c, {}),
        ("tools", "Which tools carried the session?", A.q_tools(B, cols, tables), ["2A78D6"], {"horizontal": True}),
        ("layers", "Where do failures live, and which escaped to delivery?", A.q_layers(B, cols, tables),
         ["2A78D6", "D03B3B"], {"horizontal": True}),
        ("detection", "What actually caught the failures?", A.q_detection(B, cols, tables), ["2A78D6"], {"horizontal": True}),
        ("latency", "How late were failures caught?", A.q_latency(B, cols, tables), ["2A78D6"], {}),
        ("time", "Where did the measured tool time go?", A.q_time(B, cols, tables), ["2A78D6"], {"horizontal": True}),
    ]
    plan += [(f"iter-{s}", f"Did the {s} loop converge?", A.q_iterations(B, tables, s), ["0CA30C", "D03B3B"], {"line": True})
             for s in A.iteration_sets(tables)]
    plan += [(f"how-{s}", f"How were the {s} defects found?", A.q_how_found(B, tables, s), ["2A78D6"], {"horizontal": True})
             for s in tables["sessions"][0].get("how_found_sets") or [] if any(f["set"] == s for f in tables["findings"])]
    takeaways, row = {}, 5 + len(plan) + 1
    for i, (key, question, (block, takeaway), colours, opts) in enumerate(plan):
        chart = (C.line(data_ws, block, question, colours) if opts.get("line")
                 else C.bar(data_ws, block, question, colours, horizontal=opts.get("horizontal", False)))
        toc = charts_ws.cell(row=4 + i, column=2, value=f"{i + 1}. {question}")
        toc.hyperlink, toc.font = f"#'Charts'!B{row}", S.LINK
        row = C.place(charts_ws, row, question, takeaway, chart)
        takeaways[key] = takeaway
    return takeaways


def kpis(tables, cols) -> list:
    tc, fl = "'Tool Calls'", "Failures"
    last_c, last_f = FIRST + len(tables["tool_calls"]) + 200, FIRST + len(tables["failures"]) + 200
    o = f"{tc}!${cols['c_out']}${FIRST}:${cols['c_out']}${last_c}"
    n_calls = f"COUNTA({tc}!${cols['c_id']}${FIRST}:${cols['c_id']}${last_c})"
    return [
        ("Tool calls", f"={n_calls}", "every call the record holds", "2A78D6", "0"),
        ("Result recorded", f'=1-COUNTIF({o},"unrecorded")/{n_calls}', "calls whose result anyone can check", "898781", "0%"),
        ("Calls that failed", f'=COUNTIF({o},"error")+COUNTIF({o},"error_unflagged")+COUNTIF({o},"attested_fail")',
         "flagged, unflagged or attested", "D03B3B", "0"),
        ("Missed by the flag", f'=COUNTIF({o},"error_unflagged")', "failed while the tool reported success", "EC835A", "0"),
        ("Failures logged", f"=COUNTA({fl}!${cols['f_id']}${FIRST}:${cols['f_id']}${last_f})", "every layer, not just tools",
         "EB6834", "0"),
        ("Caught after delivery", f'=COUNTIF({fl}!${cols["f_when"]}${FIRST}:${cols["f_when"]}${last_f},"*after-delivery")',
         "reached the person before anyone caught them", "D03B3B", "0"),
    ]


def render_workbook(tables: dict, out_path: Path) -> Path:
    wb = Workbook()
    wb.active.title = "Index"
    wb.create_sheet("Charts")
    pos = positions(tables)
    calls_spec = calls_tab(wb, tables["tool_calls"], pos)
    failures_spec = failures_tab(wb, tables["failures"], pos)
    cols = column_map(calls_spec, failures_spec)
    pos.update(calls_ex_col=cols["c_ex"], fail_ex_col=cols["f_ex"])
    exchanges_tab(wb, tables["exchanges"], None, pos)
    findings_tab(wb, tables["findings"])
    sources_tab(wb, tables["sources"])
    deliverables_tab(wb, tables["deliverables"])
    open_items_tab(wb, tables["open_items"])
    events_tab(wb, tables["events"])
    build_dictionary(wb)
    takeaways = build_charts(wb, tables, cols)
    counts = {name: len(tables[key]) for name, key in
              [("Exchanges", "exchanges"), ("Tool Calls", "tool_calls"), ("Failures", "failures"), ("Findings", "findings"),
               ("Sources", "sources"), ("Deliverables", "deliverables"), ("Open Items", "open_items"), ("Events", "events")]}
    headlines = [takeaways[k] for k in ("visibility", "layers", "detection", "work")]
    build_index(wb["Index"], tables, kpis(tables, cols), headlines, counts)
    wb._sheets = [wb[n] for n in ORDER] + [s for s in wb._sheets if s.title not in ORDER]
    wb.active = 0
    out_path = Path(out_path)
    out_path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(out_path)
    return out_path
